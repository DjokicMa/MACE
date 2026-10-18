"""The writers of a materials export: `mace database --action export` and the
GUI's Database view export use this one class, so both write the same bytes
in every format (csv, json, excel, latex, html).

Standard library only at import (pandas and openpyxl load for an Excel
export): the GUI imports it without the mace.database package.
"""

import csv
import io
import json
import re
from pathlib import Path
from typing import List, Dict, Optional
from datetime import datetime

# Single source of truth for the version (mace/__init__.py has no imports, so
# this is side-effect free); fall back defensively if the package is unavailable.
try:
    from mace import __version__ as MACE_VERSION
except Exception:  # pragma: no cover - defensive
    MACE_VERSION = "unknown"

FORMATS = ('csv', 'json', 'excel', 'latex', 'html')
EXTENSIONS = {'csv': 'csv', 'json': 'json', 'excel': 'xlsx', 'latex': 'latex', 'html': 'html'}


class ExportFormatter:
    """Handles formatting and exporting of materials data to various formats."""
    
    def __init__(self):
        """Initialize the formatter."""
        self.supported_formats = ['csv', 'json', 'excel', 'latex', 'html']
        
    def export(self, data: List[Dict], output_file: str, format: str = 'csv',
               properties_filter: List[str] = None,
               include_structures: bool = False,
               metadata: Dict = None) -> bool:
        """
        Export data to specified format.
        
        Args:
            data: List of records to export (materials or properties)
            output_file: Output file path
            format: Export format (csv, json, excel, latex, html)
            properties_filter: List of property names to include (None = all)
            include_structures: Whether to include structure data
            metadata: Additional metadata to include in export
            
        Returns:
            True if successful
        """
        out = self.render(data, format, properties_filter, include_structures, metadata)
        if out is None:
            return False
        output_path = Path(output_file)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_bytes(out)
        return True

    def render(self, data: List[Dict], format: str = 'csv',
               properties_filter: List[str] = None,
               include_structures: bool = False,
               metadata: Dict = None) -> Optional[bytes]:
        """The bytes export() writes (None where it writes nothing: no rows,
        except for JSON). The GUI's Database view export sends these."""
        format = format.lower()
        if format not in self.supported_formats:
            raise ValueError(f"Unsupported format: {format}. Choose from {self.supported_formats}")

        # Filter data if needed
        if properties_filter and data:
            data = self._filter_properties(data, properties_filter)

        # Remove structure data if not requested
        if not include_structures and data:
            data = self._remove_structure_data(data)

        if format == 'excel':
            return self._excel_bytes(data, metadata)
        buf = io.StringIO(newline='')
        ok = {'csv': lambda: self._write_csv(data, buf),
              'json': lambda: self._write_json(data, buf, metadata),
              'latex': lambda: self._write_latex(data, buf, metadata),
              'html': lambda: self._write_html(data, buf, metadata)}[format]()
        return buf.getvalue().encode('utf-8') if ok else None

    def _filter_properties(self, data: List[Dict], properties: List[str]) -> List[Dict]:
        """Filter data to include only specified properties (a name renamed
        in 1.1.6 by its current name, the column the export writes)."""
        from mace.utils.property_keys import canonical
        properties = {canonical(p) for p in properties}
        filtered = []
        for record in data:
            filtered_record = {}
            for key, value in record.items():
                if key in properties or key in ['material_id', 'calc_id']:  # Always include IDs
                    filtered_record[key] = value
            filtered.append(filtered_record)
        return filtered
        
    def _remove_structure_data(self, data: List[Dict]) -> List[Dict]:
        """Remove large structure-related fields."""
        from mace.utils.db_export_values import STRUCTURE_FIELDS
        structure_fields = STRUCTURE_FIELDS
        
        cleaned = []
        for record in data:
            cleaned_record = {}
            for key, value in record.items():
                if key not in structure_fields:
                    cleaned_record[key] = value
            cleaned.append(cleaned_record)
        return cleaned
        
    def _write_csv(self, data: List[Dict], f) -> bool:
        """Export to CSV format."""
        if not data:
            return False
            
        
        # Get all unique keys
        all_keys = set()
        for record in data:
            all_keys.update(record.keys())
        fieldnames = sorted(all_keys)
        
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(data)

        return True
        
    def _write_json(self, data: List[Dict], f, metadata: Dict = None) -> bool:
        """Export to JSON format."""
        
        export_data = {
            'metadata': {
                'export_date': datetime.now().isoformat(),
                'record_count': len(data),
                'mace_version': MACE_VERSION
            },
            'data': data
        }
        
        if metadata:
            export_data['metadata'].update(metadata)
            
        json.dump(export_data, f, indent=2, default=str)

        return True
        
    def _excel_bytes(self, data: List[Dict], metadata: Dict = None) -> Optional[bytes]:
        """Export to Excel format with formatting (the .xlsx file's bytes)."""
        try:
            import pandas as pd
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils.dataframe import dataframe_to_rows
        except ImportError:
            raise ImportError("Excel export needs openpyxl, an optional package (requirements-optional.txt): "
                              "pip install 'openpyxl<4', or export as CSV")
            
        if not data:
            return None

        
        # Convert to DataFrame. A cell holds one value: a dict or list (a
        # material's "metadata") goes in as its JSON text; openpyxl refused it,
        # so every export of a database with metadata failed.
        df = pd.DataFrame([{k: (json.dumps(v, default=str) if isinstance(v, (dict, list, tuple)) else v)
                            for k, v in r.items()} for r in data])
        
        # Create workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Materials Data"
        
        # Add metadata sheet if provided
        if metadata:
            meta_ws = wb.create_sheet("Metadata")
            meta_ws['A1'] = "Export Information"
            meta_ws['A1'].font = Font(bold=True, size=14)
            
            row = 3
            for key, value in metadata.items():
                meta_ws[f'A{row}'] = key
                meta_ws[f'B{row}'] = str(value)
                row += 1
                
            meta_ws['A{}'.format(row+1)] = "Export Date"
            meta_ws['B{}'.format(row+1)] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            
        # Write data to main sheet
        for r in dataframe_to_rows(df, index=False, header=True):
            ws.append(r)
            
        # Format header row
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_alignment = Alignment(horizontal="center")
        
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
            
        # Save workbook
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()
        
    def _write_latex(self, data: List[Dict], f, metadata: Dict = None) -> bool:
        """Export to LaTeX table format."""
        if not data:
            return False
            
        
        # Select key columns for LaTeX table (too many columns don't work well).
        # Use the canonical property name 'total_energy_au' (the extractor/DB never
        # emit a bare 'total_energy', so it was always silently dropped here).
        key_columns = ['material_id', 'formula', 'space_group', 'band_gap',
                      'total_energy_au', 'optimization_converged']
        
        # Filter to available columns
        available_columns = []
        if data:
            first_record = data[0]
            for col in key_columns:
                if col in first_record:
                    available_columns.append(col)
                    
        if not available_columns:
            # Fallback to first 6 columns
            available_columns = list(data[0].keys())[:6]
            
        # Write LaTeX preamble
        f.write("% MACE Database Export\n")
        f.write(f"% Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
        if metadata:
            f.write(f"% Metadata: {metadata}\n")
        f.write("\n")

        # Begin table
        f.write("\\begin{table}[htbp]\n")
        f.write("\\centering\n")
        f.write("\\caption{Materials Database Export}\n")
        f.write("\\label{tab:materials}\n")

        # Column specification
        col_spec = "|" + "|".join(["c"] * len(available_columns)) + "|"
        f.write(f"\\begin{{tabular}}{{{col_spec}}}\n")
        f.write("\\hline\n")

        # Header row
        headers = [self._latex_escape(col.replace('_', ' ').title()) 
                  for col in available_columns]
        f.write(" & ".join(headers) + " \\\\\n")
        f.write("\\hline\\hline\n")

        # Data rows
        for record in data[:50]:  # Limit to 50 rows for LaTeX
            row_data = []
            for col in available_columns:
                value = record.get(col, '')
                # Format numeric values
                if isinstance(value, float):
                    value = f"{value:.4f}"
                row_data.append(self._latex_escape(str(value)))
            f.write(" & ".join(row_data) + " \\\\\n")
            f.write("\\hline\n")

        # End table
        f.write("\\end{tabular}\n")
        f.write("\\end{table}\n")

        if len(data) > 50:
            f.write(f"\n% Note: Showing first 50 of {len(data)} total records\n")

        return True
        
    def _write_html(self, data: List[Dict], f, metadata: Dict = None) -> bool:
        """Export to HTML table format with styling."""
        if not data:
            return False
            
        
        # HTML header
        f.write("""<!DOCTYPE html>
<html>
<head>
    <title>MACE Database Export</title>
    <style>
        body { font-family: Arial, sans-serif; margin: 20px; }
        h1 { color: #366092; }
        table { border-collapse: collapse; width: 100%; margin-top: 20px; }
        th { background-color: #366092; color: white; padding: 10px; text-align: left; }
        td { border: 1px solid #ddd; padding: 8px; }
        tr:nth-child(even) { background-color: #f2f2f2; }
        tr:hover { background-color: #e8f4f8; }
        .metadata { background-color: #f0f0f0; padding: 10px; margin-bottom: 20px; }
    </style>
</head>
<body>
    <h1>MACE Database Export</h1>
""")
        
        # Metadata section
        if metadata or True:
            f.write('    <div class="metadata">\n')
            f.write('        <h3>Export Information</h3>\n')
            f.write(f'        <p><strong>Export Date:</strong> {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}</p>\n')
            f.write(f'        <p><strong>Total Records:</strong> {len(data)}</p>\n')
            if metadata:
                for key, value in metadata.items():
                    # escaped: a filter such as 'band_gap < 3' is text, not markup
                    f.write(f'        <p><strong>{self._html_escape(str(key))}:</strong> '
                            f'{self._html_escape(str(value))}</p>\n')
            f.write('    </div>\n')
            
        # Data table
        f.write('    <table>\n')
        
        # Header
        if data:
            f.write('        <tr>\n')
            for key in data[0].keys():
                header = key.replace('_', ' ').title()
                f.write(f'            <th>{self._html_escape(header)}</th>\n')
            f.write('        </tr>\n')
            
            # Data rows
            for record in data:
                f.write('        <tr>\n')
                for key in data[0].keys():
                    value = record.get(key, '')
                    if isinstance(value, float):
                        value = f"{value:.6f}".rstrip('0').rstrip('.')
                    f.write(f'            <td>{self._html_escape(str(value))}</td>\n')
                f.write('        </tr>\n')
                
        f.write('    </table>\n')
        f.write('</body>\n</html>')
        
        return True

    _LATEX_SPECIALS = {
        '&': r'\&',
        '%': r'\%',
        '$': r'\$',
        '#': r'\#',
        '_': r'\_',
        '{': r'\{',
        '}': r'\}',
        '~': r'\textasciitilde{}',
        '^': r'\^{}',
        '\\': r'\textbackslash{}',
    }
    _LATEX_RE = re.compile('|'.join(re.escape(c) for c in _LATEX_SPECIALS))

    def _latex_escape(self, text: str) -> str:
        """Escape special LaTeX characters in one pass (replacing them in turn
        re-escaped the backslashes just inserted: "1_dia" became
        "1\\textbackslash{}_dia", which does not compile)."""
        return self._LATEX_RE.sub(lambda m: self._LATEX_SPECIALS[m.group(0)], text)
        
    def _html_escape(self, text: str) -> str:
        """Escape HTML special characters."""
        replacements = {
            '&': '&amp;',
            '<': '&lt;',
            '>': '&gt;',
            '"': '&quot;',
            "'": '&#39;',
        }
        for char, replacement in replacements.items():
            text = text.replace(char, replacement)
        return text


