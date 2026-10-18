"""`mace database --action export` through the CLI.

HPCC QA (2026-10-03): a plain CSV/JSON export held only the materials table
(the CLI meant 'all properties' by include_properties=None, which the
exporter read as 'none'), and `--format excel` without openpyxl wrote CSV
text into the .xlsx file and reported success.
"""
import csv
import json
import os
import subprocess
import sys

import pytest

from conftest import REPO_ROOT

MACE_CLI = REPO_ROOT / "mace_cli"


@pytest.fixture
def workdir(tmp_path):
    sys.path.insert(0, str(REPO_ROOT))
    from mace.database.materials import MaterialDatabase
    db = MaterialDatabase(str(tmp_path / "materials.db"))
    db.create_material("MgO", "MgO", space_group=225)
    db.store_material_property("MgO", "band_gap", 7.6)
    db.store_material_property("MgO", "total_energy", -275.1)
    return tmp_path


def _env(extra_path=None):
    env = dict(os.environ)
    env["MACE_NO_BANNER"] = "1"
    for var in ("MACE_CONTEXT_DIR", "MACE_WORKFLOW_ID", "MACE_ISOLATION_MODE",
                "MACE_HOME", "PYTHONPATH"):
        env.pop(var, None)
    if extra_path:
        env["PYTHONPATH"] = str(extra_path)
    return env


def _export(cwd, *args, env=None):
    return subprocess.run(
        [sys.executable, str(MACE_CLI), "--no-banner", "database", "--action", "export", *args],
        cwd=str(cwd), env=env or _env(), capture_output=True, text=True, timeout=300)


def test_csv_export_includes_properties(workdir):
    proc = _export(workdir, "--format", "csv", "--output", "mgo.csv")
    assert proc.returncode == 0, proc.stdout + proc.stderr
    with open(workdir / "mgo.csv") as fh:
        (row,) = list(csv.DictReader(fh))
    assert row["formula"] == "MgO"
    assert float(row["band_gap"]) == pytest.approx(7.6)
    assert float(row["total_energy"]) == pytest.approx(-275.1)


def test_json_export_includes_properties(workdir):
    proc = _export(workdir, "--format", "json", "--output", "mgo.json")
    assert proc.returncode == 0, proc.stdout + proc.stderr
    text = (workdir / "mgo.json").read_text()
    assert "band_gap" in text and "total_energy" in text


def test_excel_without_openpyxl_writes_no_fake_xlsx(workdir, tmp_path):
    blocker = tmp_path / "block"
    (blocker / "openpyxl").mkdir(parents=True)
    (blocker / "openpyxl" / "__init__.py").write_text(
        "raise ImportError('openpyxl blocked for this test')\n")
    proc = _export(workdir, "--format", "excel", "--output", "mgo.xlsx",
                   env=_env(blocker))
    out = proc.stdout + proc.stderr
    assert proc.returncode != 0, out
    assert not (workdir / "mgo.xlsx").exists(), out
    assert "openpyxl" in out


def test_excel_export_of_a_material_with_metadata(tmp_path):
    """Every database a build writes has materials metadata (a dict): the
    Excel export failed on it ("Cannot convert {...} to Excel"). The cell
    now holds the dict's JSON text."""
    pytest.importorskip("openpyxl")
    pytest.importorskip("pandas")
    sys.path.insert(0, str(REPO_ROOT))
    from mace.database.materials import MaterialDatabase
    db = MaterialDatabase(str(tmp_path / "materials.db"))
    db.create_material("C", "C", space_group=227, metadata={"cell_formula": "C2"})
    proc = _export(tmp_path, "--format", "excel", "--output", "c.xlsx")
    assert proc.returncode == 0, proc.stdout + proc.stderr
    from openpyxl import load_workbook
    ws = load_workbook(tmp_path / "c.xlsx")["Materials Data"]
    header, row = [c.value for c in ws[1]], [c.value for c in ws[2]]
    assert json.loads(row[header.index("metadata")]) == {"cell_formula": "C2"}


def test_html_export_escapes_the_filter_text(workdir):
    """A filter such as 'band_gap < 9' is shown as text in the HTML export, not read as markup."""
    proc = _export(workdir, "--format", "html", "--output", "f.html", "--filter", "band_gap < 9")
    assert proc.returncode == 0, proc.stdout + proc.stderr
    html = (workdir / "f.html").read_text()
    assert "band_gap &lt; 9" in html and "band_gap < 9" not in html
