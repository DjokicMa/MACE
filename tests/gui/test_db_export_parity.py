"""The Database view's export is `mace database --action export` (CLI/GUI
parity, 1.1.6): the GUI hands its filtered, sorted rows to the CLI's own
export writer, so for one database and one filter the two give the same
bytes in every format (csv, json, excel, latex, html), up to the export time.

The database is built from real corpus outputs with the real CLI
(`mace analyze --extract-properties --register`); the CLI export runs as a
subprocess and the GUI's through its HTTP route."""
import gzip
import io
import json
import re
import shutil
import subprocess
import sys
import urllib.request
from pathlib import Path

import pytest

import guiharness as H

CORPUS = H.REPO_ROOT / "test"
RUNS = ["OPT/1_dia_opt_rev1", "SP/1_dia_opt_rev1_sp_B3LYP-D3-D3_optimized",
        "OPT/1LiFSI-1DEC-conf1_MOLECULE_OPT_symm_HSESOL3C_SOLDEF2MSVP_opt_HSESOL3C_optimized",
        "OPT/1LiFSI-1DEC-conf2_MOLECULE_OPT_symm_HSESOL3C_SOLDEF2MSVP_opt_HSESOL3C_optimized"]
# one filter, as the GUI's filter bar sends it and as the CLI's --filter writes it
FILTERS = {"all": ("", None), "molecules": ("&dim=MOLECULE", "(dimensionality = 'MOLECULE')"),
           "sg200": ("&sg_min=200", "(space_group >= 200)")}
FORMATS = {"csv": "csv", "json": "json", "excel": "xlsx", "latex": "latex", "html": "html"}


def _env(xdg):
    return {**H.env(xdg), "NO_COLOR": "1", "MACE_NO_BANNER": "1", "OMP_NUM_THREADS": "1"}


@pytest.fixture(scope="module")
def served(tmp_path_factory):
    if not all((CORPUS / (r + ".out")).is_file() for r in RUNS):
        pytest.skip("test/ data corpus not present (gitignored)")
    tmp = tmp_path_factory.mktemp("dbparity")
    root = tmp / "proj"
    runs = root / "runs"
    runs.mkdir(parents=True)
    for r in RUNS:
        for ext in (".out", ".d12"):
            shutil.copyfile(CORPUS / (r + ext), runs / (Path(r).name + ext))
    p = subprocess.run([sys.executable, str(H.MACE_CLI), "--no-banner", "analyze", "--extract-properties", str(runs),
                        "--db-path", str(root / "materials.db"), "--register"], capture_output=True, text=True,
                       env=_env(tmp / "xdg"), cwd=str(root), timeout=900, stdin=subprocess.DEVNULL)
    assert p.returncode == 0, p.stdout[-2000:] + p.stderr[-2000:]
    proc, base, _ = H.start_server(root, tmp / "xdg", MACE_GUI_NO_WARMUP="1")
    yield root, base, tmp
    proc.terminate()
    proc.wait(timeout=15)


def _gui(base, query):
    r = urllib.request.Request(base + "api/db/export?path=materials.db&sort=created_at&dir=desc" + query,
                               headers={"X-Requested-With": "mace-gui", "X-Mace-Token": H.TOKEN})
    try:
        with urllib.request.urlopen(r, timeout=300) as resp:
            raw = resp.read()
            if resp.headers.get("Content-Encoding") == "gzip":
                raw = gzip.decompress(raw)
            return resp.status, dict(resp.headers), raw
    except urllib.error.HTTPError as e:
        return e.code, dict(e.headers), e.read()


def _cli(root, tmp, fmt, flt):
    out = tmp / f"cli_{fmt}_{abs(hash(flt))}.{FORMATS[fmt]}"
    args = [sys.executable, str(H.MACE_CLI), "--no-banner", "database", "--db-path", str(root / "materials.db"),
            "--action", "export", "--format", fmt, "--output", str(out)]
    if flt:
        args += ["--filter", flt]
    p = subprocess.run(args, capture_output=True, text=True, env=_env(tmp / "xdg"), cwd=str(tmp),
                       timeout=300, stdin=subprocess.DEVNULL)
    assert out.is_file(), p.stdout[-2000:] + p.stderr[-2000:]
    return out.read_bytes()


_DATES = [re.compile(rb'"export_date": "[^"]*"'), re.compile(rb"% Generated: [^\n]*"),
          re.compile(rb"<strong>Export Date:</strong> [^<]*")]


def _no_dates(b: bytes) -> bytes:
    for rx in _DATES:
        b = rx.sub(b"DATE", b)
    return b


def _no_filter_text(b: bytes) -> bytes:
    """The filter's own words: the CLI writes its --filter text, the GUI its filter bar's."""
    b = re.sub(rb'(?s),?\s*"filters": \[.*?\](,\s*"filter_(type|logic)": "[^"]*")?', b"", b)
    b = re.sub(rb"% Metadata: [^\n]*", b"META", b)
    return re.sub(rb"<p><strong>(filters|filter_type|filter_logic):</strong> [^<]*</p>\n\s*", b"", b)


def _cells(xlsx: bytes):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xlsx))
    out = {}
    for ws in wb.worksheets:
        rows = [tuple(c.value for c in r) for r in ws.iter_rows()]
        if ws.title == "Metadata":
            rows = [r for r in rows if r and r[0] not in ("Export Date", "filters", "filter_type", "filter_logic")]
        out[ws.title] = rows
    return out


@pytest.mark.parametrize("fmt", list(FORMATS))
@pytest.mark.parametrize("name", list(FILTERS))
def test_gui_export_is_the_clis(served, fmt, name):
    root, base, tmp = served
    query, flt = FILTERS[name]
    if fmt == "excel":
        pytest.importorskip("openpyxl")
        pytest.importorskip("pandas")
    cli = _cli(root, tmp, fmt, flt)
    st, h, ours = _gui(base, query + f"&format={fmt}")
    assert st == 200, ours[:400]
    assert "attachment" in h["Content-Disposition"] and h["Content-Disposition"].endswith(f'.{FORMATS[fmt]}"')
    if fmt == "excel":
        assert _cells(ours) == _cells(cli)
        return
    a, b = _no_dates(ours), _no_dates(cli)
    if flt:
        a, b = _no_filter_text(a), _no_filter_text(b)
    if fmt == "csv":
        assert ours == cli                     # no date in a CSV: the very bytes
    assert a == b


def test_the_filtered_json_names_the_gui_filter(served):
    root, base, _ = served
    st, _, ours = _gui(base, "&dim=MOLECULE&format=json")
    meta = json.loads(ours)["metadata"]
    assert meta["export_type"] == "filtered_materials" and meta["record_count"] == 2
    assert meta["filters"] == ["dimensionality: MOLECULE"]


def test_an_unknown_format_is_refused(served):
    root, base, _ = served
    st, _, body = _gui(base, "&format=docx")
    assert st == 400, body[:300]


def test_both_exports_write_the_116_columns(served):
    """1.1.6's key clean-up: the columns are the registry's current names; no
    renamed, removed or side-table (property_blobs) name is a column, and the
    GUI's export writes the same header."""
    import csv
    from mace.utils import property_keys as PK
    root, base, tmp = served
    header = next(csv.reader(io.StringIO(_cli(root, tmp, "csv", None).decode())))
    st, _, ours = _gui(base, "&format=csv")
    assert st == 200 and next(csv.reader(io.StringIO(ours.decode()))) == header
    assert {"cpu_time_total_s", "wall_time_s", "spin_polarized", "calculation_type", "band_gap"} <= set(header)
    assert not set(header) & (set(PK.ALIASES) | PK.REMOVED | PK.BLOB_KEYS)
    props = [h for h in header if h in PK.KEYS]
    assert props, header
